# -*- coding: utf-8 -*-
"""Module DÉCHARGE : on colle (Ctrl+V) ou dépose la capture du tableau reçu du client par mail ; le site le relit cellule par cellule
(decharge_ocr), le montre dans un tableau corrigeable et rend un Excel identique à celui du client (DECHARGE.xlsx)."""

from __future__ import annotations

import hashlib
import io
import re

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

import decharge_ocr as dro

HEADERS = {"COMPTE": "COMPTE", "NGP": "NGP", "ORDRE": "POSITION", "QTE NET": "QTE NET", "QTE BRUT": "QTE BRUT", "T/O/N": "T/O/N", "TX %": "TX %"}

PASTE_JS = """
<script>
(function(){
  var w = window.parent, d = w.document;
  if (w.__timarPaste) return; w.__timarPaste = true;
  d.addEventListener('paste', function(e){
    var items = (e.clipboardData && e.clipboardData.items) || [];
    for (var i = 0; i < items.length; i++){
      if (items[i].kind === 'file' && items[i].type.indexOf('image/') === 0){
        var f = items[i].getAsFile(); if (!f) continue;
        var box = d.querySelector('div[class*="st-key-decharge_up"] input[type="file"]'); if (!box) return;
        var ext = (f.type.split('/')[1] || 'png').replace('jpeg','jpg');
        var file = new File([f], 'capture_' + Date.now() + '.' + ext, {type: f.type});
        var dt = new DataTransfer(); dt.items.add(file);
        box.files = dt.files; box.dispatchEvent(new Event('change', {bubbles: true}));
        e.preventDefault(); return;
      }
    }
  });
})();
</script>
"""


def postprocess(t: dict) -> dict:
    """Règles du tableau final : POSITION vide -> 1 ; QTE BRUT sans nombre (NA / vide) -> égale à QTE NET."""
    rows = []
    for r in t["rows"]:
        r = dict(r)
        if "ORDRE" in r and not str(r.get("ORDRE", "")).strip():
            r["ORDRE"] = "1"
        net = str(r.get("QTE NET", "")).strip()
        if "QTE BRUT" in r and net and not str(r.get("QTE BRUT", "")).replace(".", "", 1).isdigit():
            r["QTE BRUT"] = net
        rows.append(r)
    return {**t, "rows": rows}


def full_compte(row: dict, sep: str = "") -> str:
    """COMPTE + lettre saisie (ex. 40102220260002977 + K -> 40102220260002977K) : le numéro de déclaration / compte se termine par une lettre."""
    c = str(row.get("COMPTE", "")).strip()
    letter = str(row.get("LETTRE", "")).strip().upper()[:1]
    return f"{c}{sep}{letter}" if c and letter else c


def build_xlsx(tables: list[dict], sep: str = "") -> bytes:
    """Excel du client : titre surligné, en-tête bleu, lignes encadrées (un bloc par destinataire sur la feuille DECHARGE) + feuille LISTE à plat."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "DECHARGE"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    blue = PatternFill("solid", start_color="006699", end_color="006699")
    title_fill = PatternFill("solid", start_color="FFFF00", end_color="FFFF00")
    r = 2
    flat = []
    for t in tables:
        cols = [c for c in t["columns"] if c != "LETTRE"]
        c = ws.cell(r, 1, t["title"])
        c.font, c.fill = Font(bold=False, size=11), title_fill
        r += 1
        for j, name in enumerate(cols, start=1):
            h = ws.cell(r, j, HEADERS.get(name, name))
            h.font, h.fill, h.border = Font(bold=True, color="FFFFFF"), blue, box
        r += 1
        for row in t["rows"]:
            for j, name in enumerate(cols, start=1):
                v = full_compte(row, sep) if name == "COMPTE" else row.get(name, "")
                if name == "ORDRE" and str(v).isdigit():
                    v = int(v)
                if name in ("QTE NET", "QTE BRUT") and str(v).replace(".", "", 1).isdigit():
                    v = float(v) if "." in str(v) else int(v)
                cell = ws.cell(r, j, v if v != "" else None)
                cell.border = box
                if name in ("COMPTE", "NGP"):
                    cell.number_format = "@"
                if isinstance(v, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
            flat.append({"DESTINATAIRE": t["title"], **{HEADERS.get(n, n): (full_compte(row, sep) if n == "COMPTE" else row.get(n, "")) for n in cols}})
            r += 1
        r += 2
    widths = {1: 24, 2: 14, 3: 8, 4: 12, 5: 12, 6: 10, 7: 10}
    for k, wdt in widths.items():
        ws.column_dimensions[chr(64 + k)].width = wdt
    ws2 = wb.create_sheet("LISTE")
    if flat:
        keys = list(flat[0].keys())
        for j, k in enumerate(keys, start=1):
            h = ws2.cell(1, j, k)
            h.font, h.fill = Font(bold=True, color="FFFFFF"), blue
        for i, row in enumerate(flat, start=2):
            for j, k in enumerate(keys, start=1):
                v = row.get(k, "")
                if k in ("QTE NET", "QTE BRUT") and str(v).isdigit():
                    v = int(v)
                ws2.cell(i, j, v if v != "" else None)
        for j in range(1, len(keys) + 1):
            ws2.column_dimensions[chr(64 + j)].width = 22 if j == 1 else 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render():
    st.markdown("### 📋 DÉCHARGE")
    t1, t2 = st.tabs(["Tableau reçu en capture (image)", "Fiche d'imputation (PDF de BADR)"])
    with t1:
        _render_capture()
    with t2:
        _render_imputation()


def _render_imputation():
    import imputation_pdf as ip
    st.caption("Déposez la **fiche d'imputation** (PDF de BADR, « Régimes économiques en douane ») : le site lit toutes les lignes (compte à apurer + lettre, NGP, ordre, "
               "quantité, valeur, quantité nette, taux, origine…) et vous rend l'Excel.")
    ups = st.file_uploader("Fiche(s) d'imputation (PDF)", type=["pdf"], accept_multiple_files=True, key="imput_files")
    if not ups:
        return
    fiches: list[dict] = []
    for f in ups:
        try:
            got = ip.parse(f.getvalue())
        except Exception as exc:  # noqa: BLE001
            st.error(f"{f.name} : lecture impossible ({type(exc).__name__}: {exc}).")
            continue
        if not got or not any(g["rows"] for g in got):
            st.warning(f"{f.name} : aucune ligne d'imputation reconnue (le PDF doit être la fiche imprimée par BADR, en texte — pas une image scannée).")
            continue
        fiches += got
    for k, fi in enumerate(fiches):
        h = fi["header"]
        with st.expander(f"Fiche {h.get('REFERENCE', '?')} — {h.get('SOUMISSIONNAIRE', '')} · {len(fi['rows'])} ligne(s)", expanded=True):
            val = f"{h['VALEUR DECLARATION']:,.0f}".replace(",", " ") if "VALEUR DECLARATION" in h else "—"
            st.markdown(f"**Déclaration** {h.get('REFERENCE', '—')} · **Date** {h.get('DATE', '—')} · **Soumissionnaire** {h.get('SOUMISSIONNAIRE', '—')} "
                        f"· **Valeur déclaration** {val}")
            st.caption(f"{h.get('FICHE IMPUTATION', '')} · {h.get('BUREAU', '')}")
            bad = [m for ok, m in fi["checks"] if not ok]
            if bad:
                st.warning("À vérifier : " + " · ".join(bad))
            elif fi["checks"]:
                st.success(" · ".join(m for _, m in fi["checks"]))
            df = pd.DataFrame(fi["rows"], columns=ip.COLUMNS)
            st.dataframe(df, use_container_width=True, hide_index=True,
                         column_config={"DATE DU COMPTE": st.column_config.DateColumn(format="DD/MM/YYYY"),
                                        "VALEUR": st.column_config.NumberColumn(format="%.3f"),
                                        "QUANTITE": st.column_config.NumberColumn(format="%.3f"),
                                        "QUANTITE NETTE": st.column_config.NumberColumn(format="%.3f"),
                                        "TAUX": st.column_config.NumberColumn(format="%.3f")})
    if fiches:
        st.download_button("Télécharger FICHE_IMPUTATION.xlsx", data=ip.build_xlsx(fiches), file_name="FICHE_IMPUTATION.xlsx", type="primary",
                           icon=":material/download:", key="imput_dl", use_container_width=True,
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _render_capture():
    st.markdown("#### Tableau reçu du client")
    st.caption("Collez la capture du tableau reçu par mail (**Ctrl+V** sur cette page) ou déposez l'image : le site relit le tableau et vous rend l'Excel. "
               "Une capture entière de l'écran convient (le tableau est repéré par son en-tête bleu) ; plusieurs tableaux sur une capture sont lus ensemble.")
    components.html(PASTE_JS, height=0)
    with st.container(key="decharge_up"):
        files = st.file_uploader("Capture(s) du tableau (PNG, JPG, WEBP) — ou Ctrl+V", type=["png", "jpg", "jpeg", "webp", "bmp"],
                                 accept_multiple_files=True, key="decharge_files")
    cache = st.session_state.setdefault("decharge_cache", {})
    c1, c2 = st.columns([1, 2])
    sep_label = c1.selectbox("Lettre après le COMPTE", ["collée", "avec un espace"], key="dch_sep",
                             help="collée : 40102220260002977K · avec un espace : 40102220260002977 K")
    sep = "" if sep_label == "collée" else " "
    c2.caption("**Lettre** : tapez la lettre de chaque ligne dans la colonne « Lettre » du tableau (elle est ajoutée après le COMPTE dans l'Excel). "
               "**POSITION** vide = 1 · **QTE BRUT** « NA » = copie de QTE NET.")
    if not files:
        st.info("Aucune capture pour l'instant. Astuce : Windows + Maj + S, sélectionnez le tableau, puis revenez ici et faites Ctrl+V.")
        return
    tables: list[dict] = []
    seen: set[str] = set()
    for f in files:
        data = f.getvalue()
        key = hashlib.sha1(data).hexdigest()
        if key in seen:                                       # même capture déposée / collée deux fois : lue une seule fois
            st.caption(f"{f.name} : identique à une capture déjà lue — ignorée.")
            continue
        seen.add(key)
        if key not in cache:
            try:
                import ui_game
                ctx = ui_game.playing("Lecture du tableau… jouez en attendant")
            except Exception:  # noqa: BLE001
                from contextlib import nullcontext
                ctx = nullcontext()
            with st.spinner(f"Lecture de {f.name}…"), ctx:
                try:
                    cache[key] = dro.read_image(data)
                except Exception as exc:  # noqa: BLE001
                    cache[key] = []
                    st.error(f"{f.name} : lecture impossible ({type(exc).__name__}: {exc}).")
        found = cache[key]
        with st.expander(f"{f.name} — {len(found)} tableau(x) lu(s)", expanded=True):
            st.image(data, caption="Capture", use_container_width=True)
            if not found:
                st.warning("Aucun tableau reconnu : la capture doit montrer l'en-tête bleu (COMPTE · NGP · QTE NET …). "
                           "Prenez la capture un peu plus large et sans zoom réduit.")
            for k, t0 in enumerate(found):
                t = postprocess(t0)
                title = st.text_input("Destinataire", value=t["title"], key=f"dch_t_{key}_{k}")
                cols = ["COMPTE", "LETTRE"] + [c for c in t["columns"] if c != "COMPTE"]
                df = pd.DataFrame([{**r, "LETTRE": ""} for r in t["rows"]], columns=cols).fillna("")
                edited = st.data_editor(df, num_rows="dynamic", use_container_width=True, hide_index=True, key=f"dch_e_{key}_{k}",
                                        column_config={"ORDRE": st.column_config.TextColumn("POSITION", width="small"),
                                                       "LETTRE": st.column_config.TextColumn("Lettre", width="small", max_chars=1,
                                                                                            help="Lettre qui termine le n° de compte / déclaration (ex. K) : "
                                                                                                 "elle est ajoutée après le COMPTE dans l'Excel.")})
                rows = edited.fillna("").astype(str).to_dict("records")
                cur = {"title": title, "columns": cols, "rows": rows}
                issues = dro.validate({"rows": rows})
                bad_l = [i for i, r in enumerate(rows, start=1) if r.get("LETTRE", "").strip() and not re.fullmatch(r"[A-Za-z]", r["LETTRE"].strip())]
                if bad_l:
                    issues.append("lettre : une seule lettre A-Z attendue (ligne " + ", ".join(map(str, bad_l)) + ")")
                if issues:
                    st.warning("Cellules à vérifier (corrigez-les ci-dessus si besoin) : " + " · ".join(issues[:12]))
                else:
                    st.success(f"{len(rows)} ligne(s) lue(s) — formats COMPTE / NGP / quantités conformes.")
                tables.append(cur)
    if tables:
        data = build_xlsx(tables, sep)
        st.download_button("Télécharger DECHARGE.xlsx", data=data, file_name="DECHARGE.xlsx", type="primary", icon=":material/download:",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dch_dl", use_container_width=True)
