# -*- coding: utf-8 -*-
"""Fiche d'imputation (BADR, « Régimes économiques en douane ») : PDF -> tableau Excel.

La fiche liste, ligne par ligne, les comptes (acquits) apurés par une déclaration :
  NO · COMPTE A APURER (n° + lettre, ex. 411 023 2025 0014953 D, « du 24/05/2025 ») · NOMENCLATURE (NGP) · ORDRE · QUANTITE · UNITE ·
  VALEUR · QUANTITE NETTE · A · TAUX · CD · origine (« O: ES »).
Lecture du texte du PDF (pas d'OCR) ; contrôles : nombre de lignes annoncé et numérotation continue (la « VALEUR » de l'en-tête est celle de la déclaration, pas la somme des comptes)."""

from __future__ import annotations

import io
import re
from datetime import datetime

LINE_RE = re.compile(
    r"^\s*(\d{1,4})\s+(\d{3}\s\d{3}\s\d{4}\s\d{7}\s[A-Z])\s+(\d{8,10})\s+(\d+)\s+(\d[\d.]*)\s+([A-Z]{1,3})\s+(\d[\d.]*)\s+(\d[\d.]*)\s+([A-Z])\s+(\d[\d.]*)\s+(\S+)\s*$")
DATE_RE = re.compile(r"^\s*du\s+(\d{2}/\d{2}/\d{4})")
ORIG_RE = re.compile(r"^\s*O:\s*([A-Z]{2})")

COLUMNS = ["NO", "COMPTE A APURER", "COMPTE (compact)", "DATE DU COMPTE", "NOMENCLATURE", "ORDRE", "QUANTITE", "UNITE", "VALEUR",
           "QUANTITE NETTE", "A", "TAUX", "CD", "ORIGINE"]


def _f(s) -> float:
    return float(str(s).replace(" ", ""))


def pdf_text(data: bytes) -> list[str]:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return [(p.extract_text() or "") for p in pdf.pages]


def _header(text: str) -> dict:
    h: dict = {}

    def grab(label, rx, cast=str):
        m = re.search(rx, text)
        if m:
            h[label] = cast(m.group(1).strip())
    grab("REFERENCE", r"REFERENCE\s*:\s*([\d ]+[A-Z]?)\s*(?:\n|Nombre)")
    grab("DATE", r"\nDATE\s*:\s*(\d{2}/\d{2}/\d{4})")
    grab("JOURNEE DU", r"JOURNEE DU\s*:\s*(\d{2}/\d{2}/\d{4})")
    grab("HEURE", r"HEURE\s*:\s*([\d:]+)")
    grab("DUM VERSION", r"DUM Version\s*:\s*([^\n]+)")
    grab("ETAT", r"\n(Enregistrée[^\n]*|Non enregistrée[^\n]*)")
    grab("DEMANDEUR", r"DEMANDEUR\s*:\s*([^\n]+)")
    grab("BUREAU", r"BUREAU DE\s*:\s*([^\n]+)")
    m = re.search(r"FICHE IMPUTATION\s*:\s*([^\n]+)\n([^\n]*)", text)
    if m:
        second = re.sub(r"(?<=\w) (?=\w(?: |$))", "", m.group(2)).strip()          # « C A U T I O N » -> « CAUTION »
        h["FICHE IMPUTATION"] = (m.group(1).strip() + (" " + second if second and second.replace(" ", "").isupper() and len(second) < 30 else "")).strip()
    grab("REGISTRE", r"REGISTRE\s*:\s*([^\n]+)")
    grab("CODE", r"\n\s*CODE\s*:\s*([^\n]+)")
    m = re.search(r"\nNOM\s*:\s*([^\n]+)", text)
    if m:
        h["SOUMISSIONNAIRE"] = m.group(1).strip()
    m = re.search(r"VILLE\s*:\s*([^\n]+)\nDELAI", text)
    if m:
        h["VILLE"] = m.group(1).strip()
    m = re.search(r"VALEUR\s*:\s*(\d[\d.]*)", text)
    if m:
        h["VALEUR DECLARATION"] = _f(m.group(1))
    m = re.search(r"Nombre de lignes de compte imput[ée]es\s*:\s*(\d+)", text)
    if m:
        h["NOMBRE DE LIGNES"] = int(m.group(1))
    return h


def parse(data: bytes) -> list[dict]:
    """Fiches d'un PDF : [{"header": {...}, "rows": [{colonne: valeur}], "checks": [(ok, message)]}]."""
    pages = pdf_text(data)
    full = "\n".join(pages)
    refs = list(dict.fromkeys(re.findall(r"REFERENCE\s*:\s*(\d{3} \d{3} \d{4} \d{7} [A-Z])", full)))
    out = []
    chunks = [full] if len(refs) <= 1 else [
        "\n".join(p for p in pages if f"REFERENCE : {r}" in p) for r in refs]
    for chunk in chunks:
        lines = chunk.split("\n")
        rows = []
        for i, ln in enumerate(lines):
            m = LINE_RE.match(ln)
            if not m:
                continue
            dt = DATE_RE.match(lines[i + 1]) if i + 1 < len(lines) else None
            og = ORIG_RE.match(lines[i + 2]) if i + 2 < len(lines) else (ORIG_RE.match(lines[i + 1]) if i + 1 < len(lines) else None)
            if not og and i + 1 < len(lines):
                og = ORIG_RE.match(lines[i + 1])
            rows.append({"NO": int(m.group(1)), "COMPTE A APURER": m.group(2), "COMPTE (compact)": m.group(2).replace(" ", ""),
                         "DATE DU COMPTE": datetime.strptime(dt.group(1), "%d/%m/%Y").date() if dt else None,
                         "NOMENCLATURE": m.group(3), "ORDRE": int(m.group(4)), "QUANTITE": _f(m.group(5)), "UNITE": m.group(6),
                         "VALEUR": _f(m.group(7)), "QUANTITE NETTE": _f(m.group(8)), "A": m.group(9), "TAUX": _f(m.group(10)),
                         "CD": m.group(11), "ORIGINE": og.group(1) if og else ""})
        h = _header(chunk)
        checks = []
        if "NOMBRE DE LIGNES" in h:
            checks.append((len(rows) == h["NOMBRE DE LIGNES"], f"{len(rows)} ligne(s) lue(s) pour {h['NOMBRE DE LIGNES']} annoncée(s)"))
        if rows:
            nos = [r["NO"] for r in rows]
            checks.append((nos == list(range(1, len(nos) + 1)), "numérotation continue 1 … " + str(len(nos))))
        if rows or h:
            out.append({"header": h, "rows": rows, "checks": checks})
    return out


def build_xlsx(fiches: list[dict]) -> bytes:
    """Excel : feuille IMPUTATION (toutes les lignes, filtrables, ligne de totaux) + feuille ENTETE (une ligne par fiche)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "IMPUTATION"
    blue = PatternFill("solid", start_color="006699", end_color="006699")
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    multi = len(fiches) > 1
    cols = (["FICHE"] if multi else []) + COLUMNS
    for j, c in enumerate(cols, start=1):
        h = ws.cell(1, j, c)
        h.font, h.fill, h.border, h.alignment = Font(bold=True, color="FFFFFF"), blue, box, Alignment(horizontal="center")
    r = 2
    first = r
    for f in fiches:
        ref = f["header"].get("REFERENCE", "")
        for row in f["rows"]:
            vals = ([ref] if multi else []) + [row.get(c) for c in COLUMNS]
            for j, v in enumerate(vals, start=1):
                cell = ws.cell(r, j, v)
                cell.border = box
                name = cols[j - 1]
                if name in ("VALEUR", "QUANTITE", "QUANTITE NETTE", "TAUX"):
                    cell.number_format = "#,##0.000"
                if name == "DATE DU COMPTE":
                    cell.number_format = "DD/MM/YYYY"
                if name == "NOMENCLATURE":
                    cell.number_format = "@"
            r += 1
    last = r - 1
    if last >= first:
        ws.cell(r, 1, "TOTAL").font = Font(bold=True)
        for name in ("VALEUR", "QUANTITE NETTE"):
            j = cols.index(name) + 1
            L = ws.cell(1, j).column_letter
            c = ws.cell(r, j, f"=SUM({L}{first}:{L}{last})")
            c.font, c.number_format = Font(bold=True), "#,##0.000"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(cols)).column_letter}{max(last, 1)}"
    ws.freeze_panes = "A2"
    for j, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = {"COMPTE A APURER": 26, "COMPTE (compact)": 22, "NOMENCLATURE": 15, "FICHE": 28}.get(c, 14)
    ws2 = wb.create_sheet("ENTETE")
    keys = ["REFERENCE", "DATE", "JOURNEE DU", "HEURE", "DUM VERSION", "ETAT", "DEMANDEUR", "BUREAU", "FICHE IMPUTATION", "SOUMISSIONNAIRE",
            "REGISTRE", "CODE", "VILLE", "NOMBRE DE LIGNES", "VALEUR DECLARATION"]
    for j, k in enumerate(keys, start=1):
        h = ws2.cell(1, j, k)
        h.font, h.fill = Font(bold=True, color="FFFFFF"), blue
        ws2.column_dimensions[h.column_letter].width = 24
    for i, f in enumerate(fiches, start=2):
        for j, k in enumerate(keys, start=1):
            ws2.cell(i, j, f["header"].get(k))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
